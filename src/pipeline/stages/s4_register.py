"""Stage 4：结果登记（纯代码）。

用 openpyxl 追加写入 result.xlsx，每 N 条落盘一次。
保存每份文档的模型原始返回到 data/logs/。
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from pipeline.models import PipelineItem, StageResult
from pipeline.stages.base import BaseStage


class RegisterStage(BaseStage):
    name = "s4_register"

    def __init__(self, config: dict | None = None) -> None:
        super().__init__(config)
        self.output_dir = Path(self.config.get("output_dir", "data/output"))
        self.logs_dir = Path(self.config.get("logs_dir", "data/logs"))
        self.flush_every = int(self.config.get("flush_every", 10))
        self.output_path = self.output_dir / "result.xlsx"
        self._buffer: list[dict] = []
        self._ensure_dirs()

    def _ensure_dirs(self) -> None:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.logs_dir.mkdir(parents=True, exist_ok=True)

    def run(self, items: list[PipelineItem]) -> StageResult:
        result = StageResult()
        for item in items:
            row = self._to_row(item)
            self._buffer.append(row)
            self._save_raw_log(item)

            if len(self._buffer) >= self.flush_every:
                self.flush()

            item.status = "passed"
            result.passed.append(item)

        if self._buffer:
            self.flush()
        return result

    def flush(self) -> None:
        """强制落盘缓冲区。"""
        if not self._buffer:
            return
        self._append_rows(self._buffer)
        self._buffer.clear()

    def _append_rows(self, rows: list[dict]) -> None:
        if self.output_path.exists():
            wb = load_workbook(self.output_path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "results"
            ws.append(list(rows[0].keys()))

        for row in rows:
            ws.append(list(row.values()))
        wb.save(self.output_path)

    def _save_raw_log(self, item: PipelineItem) -> None:
        if not item.score:
            return
        log_name = f"{item.item_id}_{datetime.now().strftime('%Y%m%d%H%M%S')}.json"
        log_path = self.logs_dir / log_name
        log_path.write_text(
            json.dumps(item.score.raw_response, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    @staticmethod
    def _to_row(item: PipelineItem) -> dict:
        score = item.score
        return {
            "item_id": item.item_id,
            "file_name": item.file_name,
            "category": score.category if score else "",
            "total_score": score.total_score if score else "",
            "reason": score.reason if score else "",
            "assigned_group": item.assigned_group or "",
            "status": item.status,
            "review_reason": item.review_reason,
        }
