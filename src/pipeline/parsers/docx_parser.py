"""DOCX 解析器：合并段落文本与表格单元格文本。"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from pipeline.models import ParsedDocument


def parse_docx(path: str | Path) -> ParsedDocument:
    """解析 docx，合并段落与表格单元格文本。

    经验教训：只遍历 paragraphs 会丢失表格中的内容，
    必须把表格 cell 文本并入同一条文本管线。
    """
    doc = Document(str(path))

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    tables: list[list[list[str]]] = []
    table_texts: list[str] = []
    for table in doc.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)
            table_texts.extend(c for c in cells if c)
        tables.append(rows)

    raw_text = "\n".join(paragraphs + table_texts)

    fields = _extract_fields(paragraphs, table_texts)
    return ParsedDocument(
        file_path=Path(path),
        raw_text=raw_text,
        fields=fields,
        tables=tables,
    )


def _extract_fields(paragraphs: list[str], table_cells: list[str]) -> dict:
    """简易键值提取：遍历文本中 'key: value' 形式。

    此处为占位实现，仅识别明确键值对。字段提取逻辑耦合
    需在后续根据实际模板结构定制。
    """
    fields: dict[str, str] = {}
    all_lines = paragraphs + table_cells
    for line in all_lines:
        if ":" in line:
            key, _, value = line.partition(":")
            key = key.strip()
            value = value.strip()
            if key and value:
                fields[key] = value
    return fields
